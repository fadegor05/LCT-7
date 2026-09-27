"""Парсер расписания из CSV/XLSX файлов.

Чистая функция, без зависимостей от БД и DI.
"""

import csv
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class ScheduleFormatError(Exception):
    """Ошибка формата файла расписания."""

    pass


@dataclass(frozen=True)
class ParsedRow:
    """Распарсенная строка расписания."""

    name: str
    start: date | None
    end: date | None


@dataclass(frozen=True)
class ParsedIssue:
    """Ошибка при распарсивании строки."""

    row_number: int
    raw_row: str
    reason: str


@dataclass(frozen=True)
class ParseResult:
    """Результат парсирования расписания."""

    rows: list[ParsedRow]
    issues: list[ParsedIssue]
    rows_total: int


class ScheduleParser:
    """Парсер расписания из CSV/XLSX."""

    # Синонимы заголовков для колонок (нижний регистр, обрезанные)
    NAME_HEADERS = {"работа", "наименование", "название", "вид работ", "этап", "name"}
    START_HEADERS = {"начало", "дата начала", "старт", "start"}
    END_HEADERS = {"окончание", "дата окончания", "конец", "finish", "end"}

    # Синонимы, которые матчатся только во втором проходе _find_header —
    # слишком общие сами по себе (например, "№ этапа" — это номер, а не
    # название работы). См. docstring _find_header.
    WEAK_NAME_SYNONYMS = {"этап"}

    def parse(self, filename: str, content: bytes) -> ParseResult:
        """Распарсить расписание из файла.

        Args:
            filename: имя файла (для определения формата)
            content: содержимое файла в байтах

        Returns:
            ParseResult с успешно распарсенными строками и ошибками

        Raises:
            ScheduleFormatError: если формат файла не поддерживается или файл пуст
        """
        ext = Path(filename).suffix.lower()

        if ext == ".xlsx":
            return self._parse_xlsx(content)
        elif ext == ".csv":
            return self._parse_csv(content)
        else:
            raise ScheduleFormatError(f"Неподдерживаемый формат: {ext}")

    def _parse_xlsx(self, content: bytes) -> ParseResult:
        """Распарсить XLSX файл."""
        try:
            buffer = BytesIO(content)
            wb = load_workbook(buffer, read_only=True, data_only=True)
            ws = wb.active
            if ws is None:
                raise ScheduleFormatError("Файл пуст")
        except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as e:
            raise ScheduleFormatError("Не удалось прочитать XLSX-файл") from e

        # Собрать все строки и их значения
        all_rows = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row is None:
                continue
            # Проверить, содержит ли строка хотя бы одно непустое значение
            if any(cell is not None and str(cell).strip() for cell in row):
                all_rows.append((row_idx, row))

        if not all_rows:
            raise ScheduleFormatError("Файл пуст")

        located = self._locate_header(all_rows)
        if located is None:
            raise ScheduleFormatError(
                "Не найдена колонка с названием работы. Распознаваемые названия: "
                + ", ".join(sorted(self.NAME_HEADERS))
            )
        header, header_row_idx = located

        # Распарсить строки данных (все строки после заголовка)
        data_rows = all_rows[header_row_idx + 1 :]

        return self._parse_data_rows(data_rows, header)

    def _parse_csv(self, content: bytes) -> ParseResult:
        """Распарсить CSV файл."""
        # Попробовать различные кодировки
        text = None
        for encoding in ["utf-8-sig", "utf-8", "cp1251"]:
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            raise ScheduleFormatError("Не удалось декодировать CSV файл")

        # Разделить на строки
        lines = text.strip().split("\n")
        if not lines or not any(line.strip() for line in lines):
            raise ScheduleFormatError("Файл пуст")

        # Определить разделитель через csv.Sniffer
        delimiter = self._detect_delimiter(text)

        # Распарсить CSV
        reader = csv.reader(StringIO(text), delimiter=delimiter)

        # Собрать все строки
        all_rows = []
        for row_num, row in enumerate(reader, start=1):
            if row and any(cell.strip() for cell in row):
                all_rows.append((row_num, row))

        if not all_rows:
            raise ScheduleFormatError("Файл пуст")

        located = self._locate_header(all_rows)
        if located is None:
            raise ScheduleFormatError(
                "Не найдена колонка с названием работы. Распознаваемые названия: "
                + ", ".join(sorted(self.NAME_HEADERS))
            )
        header, header_row_idx = located

        # Распарсить строки данных
        data_rows = all_rows[header_row_idx + 1 :]

        return self._parse_data_rows(data_rows, header)

    def _detect_delimiter(self, text: str) -> str:
        """Определить разделитель CSV файла."""
        # Взять первые ~10 непустых строк как образец
        lines = text.split("\n")
        non_empty_lines = [line for line in lines if line.strip()][:10]
        sample = "\n".join(non_empty_lines)

        try:
            sniffer = csv.Sniffer()
            delimiter = sniffer.sniff(sample, delimiters=",;\t").delimiter
            return delimiter
        except csv.Error:
            # Fallback на точку с запятой
            return ";"

    def _find_header(self, row: list | tuple) -> dict[str, int] | None:
        """Найти индексы колонок имени, начала и окончания в строке.

        Совпадение — по подстроке, а не по точному равенству: реальные
        шаблоны используют длинные официальные формулировки («Наименование
        этапа выполнения контракта и (или)...»), которые не совпадают
        целиком ни с одним коротким синонимом, хотя явно его содержат.

        Двухпроходный поиск имени: сначала пробуем "сильные" синонимы (всё,
        кроме "этап"), и только если имя не нашлось — разрешаем "этап".
        Короткое слово "этап" само по себе часто встречается в служебных
        колонках вроде "№ этапа"/"Номер этапа" (это порядковый номер, а не
        название работы) — приоритет сильных синонимов не даёт им
        перехватить колонку раньше настоящего имени.

        При нескольких ячейках одной категории побеждает первая найденная.

        Returns:
            dict с ключами 'name', 'start', 'end' и индексами колонок,
            или None если не найдена колонка с названием
        """
        result = self._match_header_row(row, allow_weak_name=False)
        if "name" not in result:
            result = self._match_header_row(row, allow_weak_name=True)

        if "name" not in result:
            return None

        return result

    def _match_header_row(self, row: list | tuple, *, allow_weak_name: bool) -> dict[str, int]:
        """Один проход поиска шапки по строке (см. _find_header)."""
        result: dict[str, int] = {}
        name_synonyms = (
            self.NAME_HEADERS if allow_weak_name else self.NAME_HEADERS - self.WEAK_NAME_SYNONYMS
        )

        for idx, cell in enumerate(row):
            if cell is None:
                continue
            cell_str = str(cell).lower().strip()
            if not cell_str:
                continue
            if "name" not in result and any(syn in cell_str for syn in name_synonyms):
                result["name"] = idx
            elif "start" not in result and any(syn in cell_str for syn in self.START_HEADERS):
                result["start"] = idx
            elif "end" not in result and any(syn in cell_str for syn in self.END_HEADERS):
                result["end"] = idx

        return result

    def _row_has_date_value(self, row: Sequence[object]) -> bool:
        """Похожа ли строка на строку ДАННЫХ, а не на шапку.

        Настоящая шапка (в т.ч. вторая строка двухстрочной шапки) содержит
        подписи колонок, а не сами даты. Строка с хотя бы одной настоящей
        датой не может быть шапкой — используется, чтобы не спутать первую
        строку данных со второй строкой шапки, и чтобы саму строку данных
        не принять за однострочную шапку.
        """
        for cell in row:
            if cell is None:
                continue
            if isinstance(cell, (datetime, date)):
                return True
            if isinstance(cell, str) and cell.strip():
                parsed, _error = self._parse_date(cell)
                if parsed is not None:
                    return True
        return False

    def _merge_header_rows(self, row: Sequence[object], next_row: Sequence[object]) -> list[object]:
        """Объединить шапку из двух соседних строк в одну для матчинга.

        Официальные шаблоны часто разносят подзаголовки по двум строкам
        (например, объединённая ячейка "Сроки исполнения этапа" в строке N,
        а фактические "Начало"/"Окончание" — в строке N+1). Берём непустую
        ячейку из ВТОРОЙ строки, иначе — из первой: в объединённой ячейке
        Excel заголовок-категория "оседает" в верхней строке диапазона, а
        конкретные подписи колонок — в строке под ней.
        """
        length = max(len(row), len(next_row))
        merged: list[object] = []
        for i in range(length):
            first = row[i] if i < len(row) else None
            second = next_row[i] if i < len(next_row) else None
            merged.append(second if (second is not None and str(second).strip()) else first)
        return merged

    def _locate_header(
        self, all_rows: Sequence[tuple[int, Sequence[object]]]
    ) -> tuple[dict[str, int], int] | None:
        """Найти шапку в первых 10 непустых строках.

        Просматривает всё окно, а не останавливается на первом частичном
        совпадении: официальные шаблоны часто имеют преамбулу (реквизиты
        подрядчика, название объекта) перед самой шапкой, и преамбула может
        случайно дать неполное совпадение (например, одно только "name").
        Строка с датами не рассматривается как вторая строка двухстрочной
        шапки, а как самостоятельная шапка — только если сама по себе уже
        полная (Gantt/таймлайн-шапка); неполное совпадение на такой строке
        отбрасывается как вероятная строка данных.

        Для каждой позиции пробует текущую строку как есть и — если
        следующая строка не похожа на данные — объединение текущей и
        следующей строки. Из всех найденных кандидатов побеждает самый
        полный (больше найденных колонок); при равенстве длины — более
        ранний. Как только найден кандидат со всеми тремя колонками —
        возвращает его сразу.

        Returns:
            (header, header_row_idx), где header_row_idx — индекс в
            all_rows последней строки шапки (данные начинаются со
            следующей записи), либо None, если шапка не найдена.
        """
        limit = min(10, len(all_rows))
        best: tuple[dict[str, int], int] | None = None

        for idx in range(limit):
            row = all_rows[idx][1]
            row_has_date = self._row_has_date_value(row)

            candidates: list[tuple[dict[str, int] | None, int]] = []
            own_candidate = self._find_header(row)
            # Строка с реальными датами обычно строка данных, а не шапка —
            # но если она САМА ПО СЕБЕ уже полная (имя+начало+окончание),
            # оставляем её кандидатом: это Gantt/таймлайн-шапка, где дальше
            # по той же строке идут отдельные колонки-даты (по неделям,
            # по месяцам), и это не должно мешать. Неполное совпадение на
            # строке с датами почти всегда ложное срабатывание на строке
            # данных — там его не принимаем.
            if own_candidate and (not row_has_date or len(own_candidate) == 3):
                candidates.append((own_candidate, idx))

            if idx + 1 < limit:
                next_row = all_rows[idx + 1][1]
                if not self._row_has_date_value(next_row):
                    merged = self._merge_header_rows(row, next_row)
                    candidates.append((self._find_header(merged), idx + 1))

            for candidate, header_row_idx in candidates:
                if candidate and (best is None or len(candidate) > len(best[0])):
                    best = (candidate, header_row_idx)

            if best is not None and len(best[0]) == 3:
                return best

        return best

    def _parse_data_rows(
        self, data_rows: Sequence[tuple[int, Sequence[object]]], header: dict[str, int]
    ) -> ParseResult:
        """Распарсить строки данных.

        Args:
            data_rows: список (row_number, row_data) кортежей
            header: словарь индексов колонок

        Returns:
            ParseResult с распарсенными строками и ошибками
        """
        rows = []
        issues = []
        rows_total = 0

        for row_num, row_data in data_rows:
            # Проверить, пуста ли строка
            if not any(str(cell).strip() for cell in row_data if cell is not None):
                # Пустая строка пропускается, но не считается в rows_total
                continue

            rows_total += 1

            # Получить значения колонок
            name = self._get_cell(row_data, header.get("name"))
            start_str = self._get_cell(row_data, header.get("start"))
            end_str = self._get_cell(row_data, header.get("end"))

            # Обрезать название и привести к строке
            if isinstance(name, str):
                name = name.strip()
            elif isinstance(name, (datetime, date)):
                # Если это дата/время, преобразуем в строку
                name = str(name)
            else:
                name = str(name) if name is not None else ""

            # Проверить пустое название
            if not name:
                raw_row = " | ".join(str(cell) if cell is not None else "" for cell in row_data)
                issues.append(
                    ParsedIssue(row_number=row_num, raw_row=raw_row, reason="Пустое название")
                )
                continue

            # Распарсить даты
            start = None
            end = None
            error_reason = None

            if start_str:
                start, parse_error = self._parse_date(start_str)
                if parse_error:
                    error_reason = parse_error

            if end_str and not error_reason:
                end, parse_error = self._parse_date(end_str)
                if parse_error:
                    error_reason = parse_error

            # Проверить, что окончание >= начало
            if not error_reason and start is not None and end is not None and end < start:
                error_reason = "Дата окончания раньше даты начала"

            if error_reason:
                raw_row = " | ".join(str(cell) if cell is not None else "" for cell in row_data)
                issues.append(ParsedIssue(row_number=row_num, raw_row=raw_row, reason=error_reason))
                continue

            # Успешно распарсена строка
            rows.append(ParsedRow(name=name, start=start, end=end))

        return ParseResult(rows=rows, issues=issues, rows_total=rows_total)

    def _get_cell(self, row: Sequence[object], idx: int | None) -> str | datetime | date | None:
        """Получить значение ячейки по индексу."""
        if idx is None or idx >= len(row):
            return None
        cell = row[idx]
        if cell is None:
            return None
        # Возвращать datetime/date объекты как есть, чтобы _parse_date мог их обработать
        if isinstance(cell, (datetime, date)):
            return cell
        return str(cell).strip() if str(cell).strip() else None

    def _parse_date(self, date_input: str | datetime | date) -> tuple[date | None, str | None]:
        """Распарсить дату из строки или datetime объекта.

        Returns:
            (date, None) если успешно, (None, error_reason) если ошибка
        """
        # Если это datetime объект, преобразовать в date
        if isinstance(date_input, datetime):
            return date_input.date(), None
        elif isinstance(date_input, date):
            return date_input, None

        # Работать со строкой
        date_str = date_input.strip() if isinstance(date_input, str) else str(date_input).strip()

        # Попробовать различные форматы
        formats = [
            ("%d.%m.%Y", "dd.mm.yyyy"),
            ("%Y-%m-%d", "yyyy-mm-dd"),
            ("%d/%m/%Y", "dd/mm/yyyy"),
        ]

        for fmt, _fmt_name in formats:
            try:
                parsed_date = datetime.strptime(date_str, fmt).date()
                return parsed_date, None
            except ValueError:
                continue

        # Ошибка: не удалось распарсить
        return None, f"Не удалось разобрать дату: {date_str}"
