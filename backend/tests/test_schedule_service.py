"""ScheduleService: import_schedule, get_schedule, update_stage_dates, get_stage_window."""

import uuid
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.domains.schedule.models import Schedule, ScheduleStage
from app.domains.schedule.service import ScheduleService
from app.domains.site.models import Site
from app.domains.user.models import User
from app.domains.user.service import STAND_USER_ID, UserService
from app.plugins.exceptions import AppError


async def seed(container) -> None:
    """Создаёт стендового пользователя."""
    async with container() as request_container:
        service = await request_container.get(UserService)
        await service.ensure_stand_user()


def create_csv_schedule(
    rows: list[tuple[str, str | None, str | None]], delimiter: str = ";"
) -> bytes:
    """Создать CSV с графиком."""
    header = f"Название{delimiter}Начало{delimiter}Окончание\n"
    lines = [header]
    for name, start, end in rows:
        start_val = start or ""
        end_val = end or ""
        lines.append(f"{name}{delimiter}{start_val}{delimiter}{end_val}\n")
    return "".join(lines).encode("utf-8")


def create_xlsx_schedule(rows: list[tuple[str, date | None, date | None]]) -> bytes:
    """Создать XLSX с графиком."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Название"
    ws["B1"] = "Начало"
    ws["C1"] = "Окончание"

    for idx, (name, start, end) in enumerate(rows, start=2):
        ws[f"A{idx}"] = name
        if start:
            ws[f"B{idx}"] = start
        if end:
            ws[f"C{idx}"] = end

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class TestScheduleServiceImport:
    """Тесты import_schedule."""

    @pytest.mark.asyncio
    async def test_import_schedule_happy_path(self, container) -> None:
        """import_schedule() успешно парсит и сохраняет график."""
        await seed(container)

        site_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=uuid.uuid4(), full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner.id)
            session.add(site)
            await session.commit()

        # Импортируем график
        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
                ("Работа 2", "06.01.2025", "10.01.2025"),
            ]
        )

        # Используем фактического владельца
        async with container() as request_container:
            actual_owner_id = owner.id

            # Получаем сервис с реальным контейнером
            service = await request_container.get(ScheduleService)
            # Перезаписываем current_user вручную в гварде

            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": actual_owner_id})()

            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        assert report.site_id == site_id
        assert report.source_filename == "schedule.csv"
        assert report.rows_total == 2
        assert report.rows_parsed == 2
        assert report.stages_with_dates == 2
        assert report.stages_without_dates == 0
        assert len(report.stages) == 2
        assert len(report.issues) == 0
        assert report.stages[0].name == "Работа 1"
        assert report.stages[0].position == 1
        assert report.stages[0].dates_source == "file"
        assert report.stages[1].name == "Работа 2"

    @pytest.mark.asyncio
    async def test_import_schedule_with_issues(self, container) -> None:
        """import_schedule() сохраняет строки с ошибками."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        # Импортируем график с ошибками
        csv_content = create_csv_schedule(
            [
                ("", "01.01.2025", "05.01.2025"),  # Пустое название
                ("Работа 1", "01.01.2025", "05.01.2025"),  # OK
                ("Работа 2", "invalid_date", "10.01.2025"),  # Неверная дата
            ]
        )

        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            service = await request_container.get(ScheduleService)

            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()

            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        assert report.rows_total == 3
        assert report.rows_parsed == 1  # Только успешная
        assert len(report.stages) == 1
        assert len(report.issues) == 2

    @pytest.mark.asyncio
    async def test_import_schedule_format_error_raises_422(self, container) -> None:
        """import_schedule() с ошибкой формата поднимает 422."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        async with container() as request_container:
            service = await request_container.get(ScheduleService)

            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()

            with pytest.raises(AppError) as exc_info:
                await service.import_schedule(site_id, "schedule.txt", b"invalid")

        assert exc_info.value.status_code == 422

    @pytest.mark.asyncio
    async def test_import_schedule_foreign_site_raises_404(self, container) -> None:
        """import_schedule() чужого сайта поднимает 404."""
        await seed(container)

        site_id = uuid.uuid4()
        other_owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            other_owner = User(id=other_owner_id, full_name="Other Owner")
            session.add(other_owner)
            site = Site(id=site_id, name="Test Site", owner_id=other_owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule([("Работа 1", "01.01.2025", "05.01.2025")])

        async with container() as request_container:
            service = await request_container.get(ScheduleService)

            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": STAND_USER_ID})()

            with pytest.raises(AppError) as exc_info:
                await service.import_schedule(site_id, "schedule.csv", csv_content)

        assert exc_info.value.status_code == 404

    @pytest.mark.asyncio
    async def test_import_schedule_idempotent(self, container) -> None:
        """import_schedule() два раза заменяет всё (не создаёт дубли)."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content_1 = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
                ("Работа 2", "06.01.2025", "10.01.2025"),
            ]
        )

        # Первый импорт
        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report1 = await service.import_schedule(site_id, "schedule.csv", csv_content_1)

        assert report1.rows_parsed == 2

        # Второй импорт с другими данными
        csv_content_2 = create_csv_schedule(
            [
                ("Новая работа 1", "01.02.2025", "05.02.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report2 = await service.import_schedule(site_id, "schedule.csv", csv_content_2)

        assert report2.rows_parsed == 1
        # Проверяем, что в БД есть только один график для site_id (re-import заменяет, не дублирует)
        # Новый график имеет новый ID (fresh imported_at)
        assert report2.imported_at >= report1.imported_at

        # Проверяем в БД
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            schedules = await session.execute(
                __import__("sqlalchemy", fromlist=["select"])
                .select(Schedule)
                .where(Schedule.site_id == site_id)
            )
            schedule_list = schedules.scalars().all()
            assert len(schedule_list) == 1

    @pytest.mark.asyncio
    async def test_import_schedule_stages_with_and_without_dates(self, container) -> None:
        """import_schedule() различает stages_with_dates и stages_without_dates."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("С датами", "01.01.2025", "05.01.2025"),  # with dates
                ("Только начало", "01.02.2025", None),  # no end -> without dates
                ("Без дат", None, None),  # no dates -> without dates
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        assert report.stages_with_dates == 1
        assert report.stages_without_dates == 2

    @pytest.mark.asyncio
    async def test_import_schedule_creates_new_schedule_on_reimport(self, container) -> None:
        """Re-import creates a new Schedule row with fresh id and imported_at."""
        await seed(container)

        import asyncio

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content_1 = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
                ("Работа 2", "06.01.2025", "10.01.2025"),
            ]
        )

        # Первый импорт
        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report1 = await service.import_schedule(site_id, "schedule.csv", csv_content_1)

        schedule_id_1 = report1.id
        imported_at_1 = report1.imported_at

        # Подождаём немного чтобы гарантировать другую timestamp
        await asyncio.sleep(0.1)

        # Второй импорт с другими данными
        csv_content_2 = create_csv_schedule(
            [
                ("Новая работа 1", "01.02.2025", "05.02.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report2 = await service.import_schedule(site_id, "schedule.csv", csv_content_2)

        schedule_id_2 = report2.id
        imported_at_2 = report2.imported_at

        # Проверяем, что это разные Schedule (новый ID)
        assert schedule_id_2 != schedule_id_1

        # Проверяем, что imported_at обновлён (или равен, если очень быстро)
        assert imported_at_2 >= imported_at_1

        # Проверяем в БД: только один schedule
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            schedules = await session.execute(
                __import__("sqlalchemy", fromlist=["select"])
                .select(Schedule)
                .where(Schedule.site_id == site_id)
            )
            schedule_list = schedules.scalars().all()
            assert len(schedule_list) == 1
            assert schedule_list[0].id == schedule_id_2

            # Проверяем, что нет дубликатов stages
            stages = await session.execute(
                __import__("sqlalchemy", fromlist=["select"])
                .select(ScheduleStage)
                .where(ScheduleStage.schedule_id == schedule_id_2)
            )
            stage_list = stages.scalars().all()
            assert len(stage_list) == 1  # Только новый этап

    @pytest.mark.asyncio
    async def test_import_schedule_concurrent_import_raises_409(self, container) -> None:
        """import_schedule() при гонке (IntegrityError) возвращает 409."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        # Первый импорт успешен
        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            await service.import_schedule(site_id, "schedule.csv", csv_content)

        # Второй импорт: имитируем IntegrityError при commit через монкейпатч
        from sqlalchemy.exc import IntegrityError

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()

            # Монкейпатч: заставляем session.commit() вызвать IntegrityError один раз
            original_commit = service._session.commit

            call_count = 0

            async def failing_commit():
                nonlocal call_count
                call_count += 1
                if call_count == 1:
                    # Первый вызов commit() поднимает IntegrityError
                    raise IntegrityError("statement", "params", "orig")
                # После rollback, нормальный commit не случится в этом тесте
                await original_commit()

            service._session.commit = failing_commit

            csv_content2 = create_csv_schedule(
                [
                    ("Работа 2", "02.01.2025", "06.01.2025"),
                ]
            )

            with pytest.raises(AppError) as exc_info:
                await service.import_schedule(site_id, "schedule.csv", csv_content2)

        assert exc_info.value.status_code == 409
        assert "уже импортируется" in exc_info.value.message


class TestScheduleServiceGet:
    """Тесты get_schedule."""

    @pytest.mark.asyncio
    async def test_get_schedule_returns_report(self, container) -> None:
        """get_schedule() возвращает ScheduleReport."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            await service.import_schedule(site_id, "schedule.csv", csv_content)

        # Получаем график
        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.get_schedule(site_id)

        assert report.site_id == site_id
        assert report.source_filename == "schedule.csv"
        assert len(report.stages) == 1

    @pytest.mark.asyncio
    async def test_get_schedule_not_imported_raises_404(self, container) -> None:
        """get_schedule() графика, не загруженного, поднимает 404."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            with pytest.raises(AppError) as exc_info:
                await service.get_schedule(site_id)

        assert exc_info.value.status_code == 404
        assert exc_info.value.message == "График не загружен"

    @pytest.mark.asyncio
    async def test_get_schedule_foreign_site_raises_404(self, container) -> None:
        """get_schedule() чужого сайта поднимает 404."""
        await seed(container)

        site_id = uuid.uuid4()
        other_owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            other_owner = User(id=other_owner_id, full_name="Other Owner")
            session.add(other_owner)
            site = Site(id=site_id, name="Test Site", owner_id=other_owner_id)
            session.add(site)
            await session.commit()

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": STAND_USER_ID})()
            with pytest.raises(AppError) as exc_info:
                await service.get_schedule(site_id)

        assert exc_info.value.status_code == 404


class TestScheduleServiceUpdateDates:
    """Тесты update_stage_dates."""

    @pytest.mark.asyncio
    async def test_update_stage_dates_happy_path(self, container) -> None:
        """update_stage_dates() обновляет даты этапа."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Обновляем даты
        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            updated = await service.update_stage_dates(
                site_id,
                stage_id,
                StageDatesUpdate(planned_start=date(2026, 1, 1), planned_end=date(2026, 1, 31)),
            )

        assert updated.planned_start == date(2026, 1, 1)
        assert updated.planned_end == date(2026, 1, 31)
        assert updated.dates_source == "manual"

    @pytest.mark.asyncio
    async def test_update_stage_dates_partial(self, container) -> None:
        """update_stage_dates() обновляет только переданные поля."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "10.01.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Обновляем только planned_start (с датой, которая не нарушит end < start)
        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            updated = await service.update_stage_dates(
                site_id, stage_id, StageDatesUpdate(planned_start=date(2025, 1, 5))
            )

        assert updated.planned_start == date(2025, 1, 5)
        assert updated.planned_end == date(2025, 1, 10)  # Осталось как было

    @pytest.mark.asyncio
    async def test_update_stage_dates_explicit_null_clears(self, container) -> None:
        """update_stage_dates() явный null очищает дату."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Очищаем planned_end
        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            updated = await service.update_stage_dates(
                site_id, stage_id, StageDatesUpdate(planned_end=None)
            )

        assert updated.planned_start == date(2025, 1, 1)
        assert updated.planned_end is None

    @pytest.mark.asyncio
    async def test_update_stage_dates_end_before_start_raises_422(self, container) -> None:
        """update_stage_dates() если end < start, поднимает 422 и не коммитит."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Пытаемся установить end < start
        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            with pytest.raises(AppError) as exc_info:
                await service.update_stage_dates(
                    site_id,
                    stage_id,
                    StageDatesUpdate(planned_start=date(2026, 1, 10), planned_end=date(2026, 1, 1)),
                )

        assert exc_info.value.status_code == 422

        # Проверяем, что в БД не изменилось
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            stage = await session.get(ScheduleStage, stage_id)
            assert stage.planned_start == date(2025, 1, 1)
            assert stage.planned_end == date(2025, 1, 5)

    @pytest.mark.asyncio
    async def test_update_stage_dates_partial_end_only_violates_start_raises_422(
        self, container
    ) -> None:
        """Partial update with only planned_end that violates resulting end<start raises 422."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.03.2026", "10.03.2026"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Пытаемся установить только planned_end < existing start
        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            with pytest.raises(AppError) as exc_info:
                await service.update_stage_dates(
                    site_id, stage_id, StageDatesUpdate(planned_end=date(2026, 2, 1))
                )

        assert exc_info.value.status_code == 422

        # Проверяем, что в БД не изменилось
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            stage = await session.get(ScheduleStage, stage_id)
            assert stage.planned_start == date(2026, 3, 1)
            assert stage.planned_end == date(2026, 3, 10)
            assert stage.dates_source == "file"  # Не поменялось на manual

    @pytest.mark.asyncio
    async def test_update_stage_dates_partial_start_only_violates_end_raises_422(
        self, container
    ) -> None:
        """Partial update with only planned_start that violates resulting end<start raises 422."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.03.2026", "10.03.2026"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Пытаемся установить только planned_start > existing end
        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            with pytest.raises(AppError) as exc_info:
                await service.update_stage_dates(
                    site_id, stage_id, StageDatesUpdate(planned_start=date(2026, 4, 1))
                )

        assert exc_info.value.status_code == 422

        # Проверяем, что в БД не изменилось
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            stage = await session.get(ScheduleStage, stage_id)
            assert stage.planned_start == date(2026, 3, 1)
            assert stage.planned_end == date(2026, 3, 10)
            assert stage.dates_source == "file"  # Не поменялось на manual

    @pytest.mark.asyncio
    async def test_update_stage_dates_unknown_stage_raises_404(self, container) -> None:
        """update_stage_dates() неизвестного этапа поднимает 404."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            with pytest.raises(AppError) as exc_info:
                await service.update_stage_dates(
                    site_id, uuid.uuid4(), StageDatesUpdate(planned_start=date(2026, 1, 1))
                )

        assert exc_info.value.status_code == 404

    @pytest.mark.asyncio
    async def test_update_stage_dates_foreign_site_raises_404(self, container) -> None:
        """update_stage_dates() этапа чужого сайта поднимает 404."""
        await seed(container)

        site_id_1 = uuid.uuid4()
        site_id_2 = uuid.uuid4()
        owner_id_1 = uuid.uuid4()
        owner_id_2 = uuid.uuid4()

        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner1 = User(id=owner_id_1, full_name="Owner 1")
            owner2 = User(id=owner_id_2, full_name="Owner 2")
            session.add(owner1)
            session.add(owner2)
            site1 = Site(id=site_id_1, name="Site 1", owner_id=owner_id_1)
            site2 = Site(id=site_id_2, name="Site 2", owner_id=owner_id_2)
            session.add(site1)
            session.add(site2)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        # Создаём график для site2
        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id_2})()
            report = await service.import_schedule(site_id_2, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Пытаемся обновить его через site1
        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id_1})()
            with pytest.raises(AppError) as exc_info:
                await service.update_stage_dates(
                    site_id_1, stage_id, StageDatesUpdate(planned_start=date(2026, 1, 1))
                )

        assert exc_info.value.status_code == 404

    @pytest.mark.asyncio
    async def test_update_stage_dates_empty_body_no_commit(self, container) -> None:
        """PATCH с пустым телом (no fields set) не должен коммитить и не менять dates_source."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)
        original_dates_source = report.stages[0].dates_source

        # Обновляем с пустым телом (no fields set)
        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            updated = await service.update_stage_dates(
                site_id,
                stage_id,
                StageDatesUpdate(),  # Empty, no fields
            )

        # Должно вернуться то же самое без изменений
        assert updated.planned_start == date(2025, 1, 1)
        assert updated.planned_end == date(2025, 1, 5)
        assert updated.dates_source == original_dates_source  # Не должно измениться

    @pytest.mark.asyncio
    async def test_update_stage_dates_same_values_no_source_change(self, container) -> None:
        """PATCH с тем же значениями не должен менять dates_source."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Обновляем с тем же значением
        from app.domains.schedule.schemas import StageDatesUpdate

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            updated = await service.update_stage_dates(
                site_id,
                stage_id,
                StageDatesUpdate(planned_start=date(2025, 1, 1), planned_end=date(2025, 1, 5)),
            )

        # Даже если передали те же значения и они не менялись, не должно менять dates_source
        assert updated.planned_start == date(2025, 1, 1)
        assert updated.planned_end == date(2025, 1, 5)
        assert updated.dates_source == "file"  # Не должно измениться


class TestScheduleServiceGetStageWindow:
    """Тесты get_stage_window."""

    @pytest.mark.asyncio
    async def test_get_stage_window_happy_path(self, container) -> None:
        """get_stage_window() возвращает (start, end) этапа."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Получаем окно
        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            start, end = await service.get_stage_window(stage_id)

        assert start == date(2025, 1, 1)
        assert end == date(2025, 1, 5)

    @pytest.mark.asyncio
    async def test_get_stage_window_no_guard_contract(self, container) -> None:
        """get_stage_window() работает без гварда (контракт для других доменов)."""
        await seed(container)

        site_id = uuid.uuid4()
        owner_id = uuid.uuid4()
        async with container() as request_container:
            session = await request_container.get(AsyncSession)
            owner = User(id=owner_id, full_name="Owner")
            session.add(owner)
            site = Site(id=site_id, name="Test Site", owner_id=owner_id)
            session.add(site)
            await session.commit()

        csv_content = create_csv_schedule(
            [
                ("Работа 1", "01.01.2025", "05.01.2025"),
            ]
        )

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            guard = service._guard
            guard._current_user = type("obj", (object,), {"id": owner_id})()
            report = await service.import_schedule(site_id, "schedule.csv", csv_content)

        stage_id = uuid.UUID(report.stages[0].id)

        # Получаем окно с другим пользователем (без гварда он должен работать)
        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            # Не устанавливаем guard._current_user или устанавливаем на другого
            start, end = await service.get_stage_window(stage_id)

        assert start == date(2025, 1, 1)
        assert end == date(2025, 1, 5)

    @pytest.mark.asyncio
    async def test_get_stage_window_unknown_stage_raises_404(self, container) -> None:
        """get_stage_window() неизвестного этапа поднимает 404."""
        await seed(container)

        async with container() as request_container:
            service = await request_container.get(ScheduleService)
            with pytest.raises(AppError) as exc_info:
                await service.get_stage_window(uuid.uuid4())

        assert exc_info.value.status_code == 404
